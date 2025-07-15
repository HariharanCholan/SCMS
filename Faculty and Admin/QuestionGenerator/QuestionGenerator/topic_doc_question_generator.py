import sys
import json
import traceback
from datetime import datetime
from bson.objectid import ObjectId
from pymongo import MongoClient
from openai import OpenAI

# ===== CONFIGURATION =====
OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
MODEL = "deepseek/deepseek-chat"
OPENROUTER_API_KEY = "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"
MONGO_URI = "mongodb+srv://HARIHARAN:Hari0519@scms.wnt7tx1.mongodb.net/?retryWrites=true&w=majority&appName=SCMS"

# ===== INITIALIZATION =====
def initialize_client():
    if not OPENROUTER_API_KEY:
        raise RuntimeError("OPENROUTER_API_KEY not set")
    return OpenAI(base_url=OPENROUTER_BASE_URL, api_key=OPENROUTER_API_KEY)

def fetch_text_from_mongo(doc_id):
    client = MongoClient(MONGO_URI)
    db = client['document_db']
    collection = db['documents']
    doc = collection.find_one({'_id': ObjectId(doc_id)})
    client.close()
    if doc and 'textContent' in doc:
        return doc['textContent'], doc.get('name', 'Untitled Document')
    raise ValueError("Document not found or missing textContent field")

def create_mixed_prompt(text_chunk, num_questions=6):
    prompt = f"""You are an exam question generator.

Generate {num_questions} questions using a mix of:
- MCQs (Single correct)
- MSQs (Multiple correct answers)
- Fill in the Blank

Format the output exactly as:

Q1 (MCQ): <question>
A) option A
B) option B
C) option C
D) option D
ANSWER: <one correct option letter>

Q2 (MSQ): <question>
A) option A
B) option B
C) option C
D) option D
ANSWER: <comma-separated correct letters like A,C>

Q3 (FILL): <Fill-in-the-blank sentence with a missing word>
ANSWER: <Correct word>

Text:
\"\"\" 
{text_chunk[:3000]}
\"\"\"
"""
    return prompt

def generate_questions(client, text_chunk, num_questions=6):
    prompt = create_mixed_prompt(text_chunk, num_questions)
    completion = client.chat.completions.create(
        extra_headers={
            "HTTP-Referer": "https://mcq-generator.com",
            "X-Title": "Mixed Question Generator"
        },
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.7,
        max_tokens=2048
    )
    return completion.choices[0].message.content

def parse_questions(raw_output, doc_name):
    questions = []
    current_question = {}
    lines = raw_output.strip().splitlines()
    
    for line in lines:
        line = line.strip()
        if line.startswith("Q"):
            if current_question:
                questions.append(current_question)
                current_question = {}
            parts = line.split(":", 1)
            if len(parts) == 2:
                qtype = "MCQ" if "(MCQ)" in parts[0] else "MSQ" if "(MSQ)" in parts[0] else "FILL"
                current_question = {
                    "type": qtype,
                    "text": parts[1].strip(),
                    "options": [],
                    "answer": "",
                    "docName": doc_name
                }
        elif line.startswith("A)") or line.startswith("B)") or line.startswith("C)") or line.startswith("D)"):
            if "options" in current_question:
                current_question["options"].append(line[2:].strip())
        elif line.startswith("ANSWER:"):
            current_question["answer"] = line.replace("ANSWER:", "").strip()

    if current_question:
        questions.append(current_question)
    
    return questions

def save_to_mongo(parsed_questions, doc_id):
    client = MongoClient(MONGO_URI)
    db = client['document_db']
    collection = db['generated_questions']
    
    result = collection.insert_one({
        "docId": ObjectId(doc_id),
        "questions": parsed_questions,
        "createdAt": datetime.utcnow()
    })
    client.close()
    return str(result.inserted_id)

def main():
    try:
        client = initialize_client()

        input_data = json.load(sys.stdin)
        document_id = input_data.get('documentId')
        num_questions = input_data.get('numQuestions', 6)

        if not document_id:
            raise ValueError("Missing documentId")

        print("Fetching text from MongoDB...")
        text, doc_name = fetch_text_from_mongo(document_id)

        if len(text) < 50:
            raise ValueError("Insufficient text in document")

        print(f"Generating {num_questions} mixed questions...")
        raw_questions = generate_questions(client, text, num_questions)

        print("Parsing questions...")
        parsed_questions = parse_questions(raw_questions, doc_name)

        print("Saving to MongoDB...")
        record_id = save_to_mongo(parsed_questions, document_id)

        output = {
            "recordId": record_id,
            "questions": parsed_questions
        }
        print(json.dumps(output))

    except Exception:
        traceback.print_exc()
        print(json.dumps({"error": "Internal error occurred"}))

if __name__ == "__main__":
    main()


'''
import os
import sys
import random
from openai import OpenAI
from PyPDF2 import PdfReader
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# ===== CONFIGURATION =====
OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
MODEL = "deepseek/deepseek-chat"  # You can replace with another model on OpenRouter
OPENROUTER_API_KEY = "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"

# ===== INITIALIZATION =====
def initialize_client():
    if not OPENROUTER_API_KEY:
        print("ERROR: OPENROUTER_API_KEY not set")
        sys.exit(1)
    return OpenAI(base_url=OPENROUTER_BASE_URL, api_key=OPENROUTER_API_KEY)

# ===== DOCUMENT PROCESSING =====
def extract_text(file_path):
    try:
        if file_path.lower().endswith('.pdf'):
            with open(file_path, 'rb') as f:
                text = "\n".join(
                    page.extract_text() for page in PdfReader(f).pages if page.extract_text()
                )
        elif file_path.lower().endswith('.docx'):
            doc = Document(file_path)
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            raise ValueError("Only PDF and DOCX files supported.")
        if len(text) < 50:
            raise ValueError("Insufficient text extracted.")
        return text
    except Exception as e:
        raise RuntimeError(f"Text extraction failed: {str(e)}")

# ===== PROMPT GENERATION =====
def create_mixed_prompt(text_chunk, num_questions=6):
    prompt = f"""You are an exam question generator.

Generate {num_questions} questions using a mix of:
- MCQs (Single correct)
- MSQs (Multiple correct answers)
- Fill in the Blank

Format the output exactly as:

Q1 (MCQ): <question>
A) option A
B) option B
C) option C
D) option D
ANSWER: <one correct option letter>

Q2 (MSQ): <question>
A) option A
B) option B
C) option C
D) option D
ANSWER: <comma-separated correct letters like A,C>

Q3 (FILL): <Fill-in-the-blank sentence with a missing word>
ANSWER: <Correct word>

Text:
\"\"\"
{text_chunk[:3000]}
\"\"\"
"""
    return prompt

# ===== QUESTION GENERATION =====
def generate_questions(client, text_chunk, num_questions=6):
    prompt = create_mixed_prompt(text_chunk, num_questions)
    try:
        completion = client.chat.completions.create(
            extra_headers={
                "HTTP-Referer": "https://mcq-generator.com",
                "X-Title": "Mixed Question Generator"
            },
            model=MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
            max_tokens=2048
        )
        return completion.choices[0].message.content
    except Exception as e:
        raise RuntimeError(f"API request failed: {str(e)}")

# ===== OUTPUT GENERATION =====
def save_to_docx(content, source_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    doc.add_heading('Generated Exam Questions', 0).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for block in content.strip().split('\n\n'):
        lines = [line.strip() for line in block.split('\n') if line.strip()]
        if lines:
            if lines[0].startswith('Q'):
                doc.add_paragraph(lines[0], style='Heading 2')
            for line in lines[1:]:
                p = doc.add_paragraph(line)
                if line.startswith('ANSWER:'):
                    p.runs[0].bold = True
    output_path = f"{os.path.splitext(source_path)[0]}_MIXED_QUESTIONS.docx"
    doc.save(output_path)
    return output_path

# ===== MAIN WORKFLOW =====
def main():
    print("=== Smart Question Generator (MCQ/MSQ/FILL) ===")
    try:
        client = initialize_client()
        doc_path = input("Enter path to your document (PDF/DOCX): ").strip('"')
        if not os.path.exists(doc_path):
            raise FileNotFoundError("File not found")

        num_q = int(input("How many questions do you want to generate? (e.g., 5, 10): "))
        if num_q < 1 or num_q > 50:
            raise ValueError("Please choose a number between 1 and 50.")

        print("Extracting text...", end=' ')
        text = extract_text(doc_path)
        print(f"✓ ({len(text.split())} words)")

        print(f"Generating {num_q} mixed questions...")
        questions = generate_questions(client, text, num_questions=num_q)

        output_path = save_to_docx(questions, doc_path)
        print(f"\n✅ Questions saved to: {os.path.abspath(output_path)}")

    except Exception as e:
        print(f"\n❌ ERROR: {str(e)}")
        if "API request failed" in str(e):
            print("- Check your OpenRouter API key")
            print("- Ensure you have credits")
            print("- Try using fewer questions or a smaller document")

if __name__ == "__main__":
    main()
'''